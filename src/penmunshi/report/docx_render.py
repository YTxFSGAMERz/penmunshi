"""Agency-grade DOCX report renderer with cover page, TOC, running headers/footers, and shaded tables."""

from __future__ import annotations

import base64
import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

from ..models import Report
from .common import count_severities, label, sort_findings

_SEV_COLORS = {
    "Critical": RGBColor(0x9B, 0x12, 0x39),
    "High": RGBColor(0xC2, 0x41, 0x0C),
    "Medium": RGBColor(0xA1, 0x62, 0x07),
    "Low": RGBColor(0x1D, 0x5F, 0x8A),
    "None": RGBColor(0x5B, 0x67, 0x73),
    "Unrated": RGBColor(0x5B, 0x67, 0x73),
}

_SEV_BG_HEX = {
    "Critical": "9B1239",
    "High": "C2410C",
    "Medium": "A16207",
    "Low": "1D5F8A",
    "None": "5B6773",
    "Unrated": "5B6773",
}

_WASH_HEX = "F4F6F3"
_MUTED_RGB = RGBColor(0x5B, 0x67, 0x73)


def _set_cell_shading(cell, hex_color: str) -> None:
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def _add_field(run, instr: str) -> None:
    escaped_instr = instr.replace('"', "&quot;")
    fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="{escaped_instr}"/>')
    run._r.append(fld)


def _mono(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8)


def _kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for key, value in rows:
        cells = table.add_row().cells
        _set_cell_shading(cells[0], _WASH_HEX)
        cells[0].text = key
        cells[1].text = value
        for run in cells[0].paragraphs[0].runs:
            run.bold = True
            run.font.name = "Segoe UI"
            run.font.size = Pt(9.5)
        for run in cells[1].paragraphs[0].runs:
            run.font.name = "Segoe UI"
            run.font.size = Pt(9.5)


def render_docx(report: Report, path: str | Path | None = None) -> bytes:
    e = report.engagement
    findings = sort_findings(report.findings)
    counts = report.severity_counts or count_severities(findings)

    doc = Document()

    # Section & Page Setup
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.different_first_page_header_footer = True

    # Running Header (pages 2+)
    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    h_run = header_p.add_run(f"PenMunshi Security Assessment | {e.project}")
    h_run.font.name = "Segoe UI"
    h_run.font.size = Pt(8.5)
    h_run.font.color.rgb = _MUTED_RGB

    # Running Footer (pages 2+)
    footer = section.footer
    footer_p = footer.paragraphs[0]
    f_left = footer_p.add_run(f"CONFIDENTIAL - {e.client}\t\tPage ")
    f_left.font.name = "Segoe UI"
    f_left.font.size = Pt(8.5)
    f_left.font.color.rgb = _MUTED_RGB

    f_page = footer_p.add_run()
    _add_field(f_page, "PAGE")
    f_page.font.name = "Segoe UI"
    f_page.font.size = Pt(8.5)
    f_page.font.color.rgb = _MUTED_RGB

    f_mid = footer_p.add_run(" of ")
    f_mid.font.name = "Segoe UI"
    f_mid.font.size = Pt(8.5)
    f_mid.font.color.rgb = _MUTED_RGB

    f_total = footer_p.add_run()
    _add_field(f_total, "NUMPAGES")
    f_total.font.name = "Segoe UI"
    f_total.font.size = Pt(8.5)
    f_total.font.color.rgb = _MUTED_RGB

    # ==================== COVER PAGE ====================
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(6)
    t_run = title_p.add_run("Penetration Test Report")
    t_run.font.name = "Georgia"
    t_run.font.size = Pt(26)
    t_run.bold = True
    t_run.font.color.rgb = RGBColor(0x18, 0x21, 0x2B)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(28)
    sub_run = sub_p.add_run(e.project)
    sub_run.font.name = "Segoe UI"
    sub_run.font.size = Pt(14)
    sub_run.font.color.rgb = _MUTED_RGB

    _kv_table(
        doc,
        [
            ("Client", e.client),
            ("Tester", e.tester or "-"),
            ("Dates", e.date_range or "-"),
            ("Methodology", e.methodology),
            ("Scope", ", ".join(e.scope) if e.scope else "-"),
            ("Reviewed by", e.reviewed_by or "(pending human review)"),
        ],
    )

    doc.add_paragraph()
    banner_p = doc.add_paragraph()
    banner_p.paragraph_format.space_before = Pt(14)
    b_run = banner_p.add_run(
        "AI-assisted draft. Scores, identifiers, quotes and secret masking were checked by "
        "software; a human tester must review every finding before delivery."
    )
    b_run.italic = True
    b_run.font.name = "Segoe UI"
    b_run.font.size = Pt(9.5)
    b_run.font.color.rgb = RGBColor(0x7A, 0x51, 0x00)

    # ==================== TABLE OF CONTENTS ====================
    doc.add_page_break()
    toc_heading = doc.add_heading("Table of contents", level=1)
    toc_heading.paragraph_format.space_after = Pt(12)
    toc_p = doc.add_paragraph()
    toc_run = toc_p.add_run()
    _add_field(toc_run, r'TOC \o "1-3" \h \z \u')
    toc_note = doc.add_paragraph()
    t_hint = toc_note.add_run("(In Microsoft Word or LibreOffice, right-click and select 'Update Field' to refresh TOC)")
    t_hint.italic = True
    t_hint.font.size = Pt(8.5)
    t_hint.font.color.rgb = _MUTED_RGB

    # ==================== EXECUTIVE SUMMARY ====================
    doc.add_page_break()
    doc.add_heading("Executive summary", level=1)
    for para in (report.executive_summary or "No summary generated.").split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())

    doc.add_heading("Severity summary", level=2)
    sum_tbl = doc.add_table(rows=1, cols=2)
    sum_tbl.style = "Table Grid"
    hdr_cells = sum_tbl.rows[0].cells
    hdr_cells[0].text = "Severity"
    hdr_cells[1].text = "Findings Count"
    for c in hdr_cells:
        _set_cell_shading(c, _WASH_HEX)
        c.paragraphs[0].runs[0].bold = True

    for k, v in counts.items():
        row_cells = sum_tbl.add_row().cells
        sev_name = label(k)
        row_cells[0].text = sev_name
        row_cells[1].text = str(v)

        # Style severity cell with colored background chip
        bg_hex = _SEV_BG_HEX.get(k, "5B6773")
        _set_cell_shading(row_cells[0], bg_hex)
        r0 = row_cells[0].paragraphs[0].runs[0]
        r0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r0.bold = True
        row_cells[1].paragraphs[0].runs[0].bold = True

    # ==================== FINDINGS OVERVIEW ====================
    doc.add_heading("Findings overview", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, head in zip(table.rows[0].cells, ["ID", "Title", "Severity", "CVSS"], strict=True):
        _set_cell_shading(cell, _WASH_HEX)
        cell.text = head
        cell.paragraphs[0].runs[0].bold = True

    for f in findings:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = f.id, f.title
        cells[2].text = label(f.severity)
        cells[3].text = f"{f.cvss_score:.1f}" if f.cvss_score is not None else "-"

        bg_hex = _SEV_BG_HEX.get(f.severity, "5B6773")
        _set_cell_shading(cells[2], bg_hex)
        r_sev = cells[2].paragraphs[0].runs[0]
        r_sev.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r_sev.bold = True

    # ==================== INDIVIDUAL FINDINGS ====================
    for f in findings:
        doc.add_page_break()
        heading = doc.add_heading(f"{f.id}: {f.title}", level=1)
        color = _SEV_COLORS.get(f.severity)
        if color is not None:
            for run in heading.runs:
                run.font.color.rgb = color
        _kv_table(
            doc,
            [
                ("Severity", f"{label(f.severity)} ({f.cvss_score if f.cvss_score is not None else 'n/a'})"),
                ("CVSS vector", f"{f.cvss_vector or 'n/a'} (computed by software)"),
                ("Affected asset", f.affected_asset or "-"),
                ("Weakness", "; ".join(f"{c.id} {c.name}" for c in f.cwe) or "-"),
                ("CVE", ", ".join(c.id for c in f.cve) or "-"),
                ("Confidence", f.confidence.replace("_", " ")),
            ],
        )
        for title, body in (("Summary", f.summary), ("Description", f.description)):
            doc.add_heading(title, level=2)
            doc.add_paragraph(body)
        doc.add_heading("Steps to reproduce", level=2)
        for i, step in enumerate(f.steps_to_reproduce, 1):
            doc.add_paragraph(f"{i}. {step}")
        doc.add_heading("Impact", level=2)
        doc.add_paragraph(f.impact)
        doc.add_heading("Remediation", level=2)
        for item in f.remediation:
            doc.add_paragraph(item, style="List Bullet")

        shown = [ev for ev in f.evidence if ev.transcript.strip() or ev.redacted_image_data_uri]
        if shown:
            doc.add_heading("Evidence", level=2)
            for ev in shown:
                doc.add_paragraph(ev.name).runs[0].bold = True
                if ev.redacted_image_data_uri and "," in ev.redacted_image_data_uri:
                    try:
                        _, b64data = ev.redacted_image_data_uri.split(",", 1)
                        img_bytes = base64.b64decode(b64data)
                        doc.add_picture(io.BytesIO(img_bytes), width=Inches(6.0))
                    except Exception:
                        pass
                if ev.transcript.strip():
                    _mono(doc, ev.transcript.strip())

    # ==================== VERIFICATION APPENDIX ====================
    doc.add_page_break()
    doc.add_heading("Appendix: verification log", level=1)
    for f in findings:
        doc.add_heading(f.id, level=2)
        ok = sum(c.supported for c in f.claims)
        doc.add_paragraph(f"Claims verified against evidence: {ok} of {len(f.claims)}", style="List Bullet")
        for r in f.removed_ids:
            doc.add_paragraph(f"Removed {r.kind.upper()} {r.id}: {r.reason}", style="List Bullet")
        for s in f.secrets:
            doc.add_paragraph(f"Masked {s.kind} {s.preview} in {s.source}", style="List Bullet")
        for i in f.lint:
            doc.add_paragraph(f"{i.level.upper()} {i.code}: {i.message}", style="List Bullet")
    t = report.tool
    doc.add_paragraph(
        f"Generated locally by PenMunshi {t.get('version', '')} with model {t.get('model', '')} on {report.generated_at}."
    )

    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    if path:
        Path(path).write_bytes(data)
    return data
