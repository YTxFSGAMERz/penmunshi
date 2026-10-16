import io

from docx import Document

from penmunshi.llm import FakeLLM
from penmunshi.models import Engagement
from penmunshi.pipeline import FindingInput, analyze_finding, build_report
from penmunshi.report import render_docx, render_html, render_markdown

NOTES = "XSS in search.\npassword=hunter2xyz\nlogs in as admin"
EVIDENCE = '<script>alert("pwned")</script> AKIAIOSFODNN7EXAMPLE'


def make_report():
    llm = FakeLLM()
    inp = FindingInput(id="F-1", notes=NOTES, pre_transcribed={"resp.txt": EVIDENCE})
    f = analyze_finding(inp, llm)
    return build_report(Engagement(client="ACME", project="Shop VAPT", tester="Tester"), [f], llm)


def test_html_escapes_hostile_evidence_and_hides_secrets():
    html = render_html(make_report())
    assert "<script>alert" not in html and "&lt;script&gt;" in html
    assert "hunter2xyz" not in html and "AKIAIOSFODNN7EXAMPLE" not in html
    assert "ACME" in html and "Appendix: verification log" in html


def test_markdown_has_sections_and_no_secrets():
    md = render_markdown(make_report())
    assert "## Executive summary" in md and "## Appendix: verification log" in md
    assert "hunter2xyz" not in md and "[REDACTED:" in md


def test_docx_roundtrip_has_text_and_no_secrets():
    data = render_docx(make_report())
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Penetration Test Report" in text and "hunter2xyz" not in text
    assert any("ACME" in cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    # Check TOC heading exists
    assert "Table of contents" in text
    # Check header has project name
    hdr_text = "\n".join(p.text for s in doc.sections for p in s.header.paragraphs)
    assert "Shop VAPT" in hdr_text
    # Check footer has client name
    ftr_text = "\n".join(p.text for s in doc.sections for p in s.footer.paragraphs)
    assert "ACME" in ftr_text
    # Check table shading was applied to at least one cell
    xml_str = doc._part.blob.decode("utf-8", errors="ignore")
    assert "w:shd" in xml_str


def test_html_print_css_and_cover_page():
    html = render_html(make_report())
    assert "@media print" in html
    assert "size: A4" in html
    assert "class=\"cover-page\"" in html
    assert "page-break-before: always" in html or "break-before: page" in html


def test_severity_counts_are_computed_by_code():
    rep = make_report()
    assert rep.severity_counts["Medium"] == 1 and rep.severity_counts["Critical"] == 0
